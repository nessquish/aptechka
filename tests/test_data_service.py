"""Экспорт, импорт и сброс данных; общие настройки уведомлений."""

import csv
import json
import tempfile
import unittest
from pathlib import Path

from pharmacy.errors import ValidationError
from pharmacy.services.container import build_services
from pharmacy.services.notification_service import NotifyOptions
from pharmacy.ui import scaling
from tests.helpers import DatabaseTestCase


class DataTestCase(DatabaseTestCase):
    def setUp(self) -> None:
        super().setUp()
        self.services = build_services(self.db)
        self.user_id = self.add_user("anna")
        self.add_product(self.user_id, name="Ибупрофен")
        self.add_product(self.user_id, name="Бинт")
        self._dir = tempfile.TemporaryDirectory()
        self.addCleanup(self._dir.cleanup)
        self.folder = Path(self._dir.name)

    def products(self, user_id=None):
        return self.services.products.list_products(user_id or self.user_id)


class ExportTest(DataTestCase):
    def test_json_has_every_product(self):
        path = self.folder / "a.json"
        self.assertEqual(
            self.services.data.export_products(self.user_id, path, "json"), 2
        )
        names = {row["name"] for row in json.loads(path.read_text(encoding="utf-8"))}
        self.assertEqual(names, {"Ибупрофен", "Бинт"})

    def test_csv_opens_in_excel_with_a_header(self):
        path = self.folder / "a.csv"
        self.services.data.export_products(self.user_id, path, "csv")
        with path.open(encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle, delimiter=";"))
        self.assertEqual(rows[0][0], "Название")
        self.assertEqual(len(rows), 3)

    def test_excel_file_is_written(self):
        from openpyxl import load_workbook

        path = self.folder / "a.xlsx"
        self.services.data.export_products(self.user_id, path, "xlsx")
        sheet = load_workbook(path).active
        self.assertEqual(sheet.max_row, 3)

    def test_unknown_format_is_rejected(self):
        with self.assertRaises(ValidationError):
            self.services.data.export_products(self.user_id, self.folder / "a", "pdf")

    def test_only_own_products_are_exported(self):
        other = self.add_user("boris")
        self.add_product(other, name="Чужой")
        path = self.folder / "a.json"
        self.services.data.export_products(self.user_id, path, "json")
        self.assertNotIn("Чужой", path.read_text(encoding="utf-8"))


class ImportTest(DataTestCase):
    def test_round_trip_through_csv(self):
        path = self.folder / "a.csv"
        self.services.data.export_products(self.user_id, path, "csv")
        other = self.add_user("boris")
        result = self.services.data.import_products(other, path)
        self.assertEqual((result.added, result.skipped), (2, 0))
        self.assertEqual(len(self.products(other)), 2)

    def test_round_trip_through_json(self):
        path = self.folder / "a.json"
        self.services.data.export_products(self.user_id, path, "json")
        other = self.add_user("boris")
        self.assertEqual(self.services.data.import_products(other, path).added, 2)

    def test_bad_rows_are_skipped_and_reported(self):
        path = self.folder / "a.json"
        path.write_text(
            json.dumps(
                [
                    {"name": "Хороший", "category": "Лекарства", "quantity": 1},
                    {"name": "", "category": "Лекарства", "quantity": 1},
                    {"name": "Без категории", "category": "Нет такой", "quantity": 1},
                ]
            ),
            encoding="utf-8",
        )
        result = self.services.data.import_products(self.user_id, path)
        self.assertEqual((result.added, result.skipped), (1, 2))
        self.assertTrue(result.errors[0].startswith("Строка 2"))

    def test_other_extensions_and_broken_files_are_rejected(self):
        for name, text in (("a.txt", "x"), ("a.json", "{oops")):
            path = self.folder / name
            path.write_text(text, encoding="utf-8")
            with self.assertRaises(ValidationError):
                self.services.data.import_products(self.user_id, path)


class ResetTest(DataTestCase):
    def test_reset_clears_the_kit_but_not_other_users(self):
        other = self.add_user("boris")
        self.add_product(other, name="Чужой")
        self.services.shopping.add_manual(self.user_id, "Вата")
        self.services.data.reset_all(self.user_id)
        self.assertEqual(self.products(), [])
        self.assertEqual(self.services.shopping.list_items(self.user_id), [])
        self.assertEqual(self.services.history.list_history(self.user_id), [])
        self.assertEqual(len(self.products(other)), 1)

    def test_deleting_the_account_removes_the_user_and_the_data(self):
        self.services.auth.delete_account(self.user_id)
        self.assertEqual(self.count_rows("products"), 0)
        self.assertEqual(self.count_rows("users"), 0)


class NotifyOptionsTest(DataTestCase):
    def kinds(self):
        return {
            n.kind for n in self.services.notifications.list_notifications(self.user_id)
        }

    def test_disabled_options_remove_everything(self):
        self.services.notifications.options_for = lambda _id: NotifyOptions(
            enabled=False
        )
        self.services.notifications.refresh(self.user_id)
        self.assertEqual(self.kinds(), set())

    def test_threshold_makes_low_stock(self):
        self.services.notifications.options_for = lambda _id: NotifyOptions(
            low_threshold=10**6
        )
        self.services.notifications.refresh(self.user_id)
        self.assertIn("low_stock", self.kinds())


class ManualScaleTest(unittest.TestCase):
    def test_manual_value_is_applied_within_limits(self):
        import os

        saved = os.environ.pop(scaling.ENV_NAME, None)
        try:
            self.assertEqual(scaling.apply_scale_factor(1.25), 1.25)
            os.environ.pop(scaling.ENV_NAME)
            self.assertEqual(scaling.apply_scale_factor(3.0), scaling.MAX_MANUAL)
            os.environ.pop(scaling.ENV_NAME)
            self.assertEqual(scaling.apply_scale_factor(0.1), scaling.MIN_MANUAL)
        finally:
            os.environ.pop(scaling.ENV_NAME, None)
            if saved is not None:
                os.environ[scaling.ENV_NAME] = saved


if __name__ == "__main__":
    unittest.main()
